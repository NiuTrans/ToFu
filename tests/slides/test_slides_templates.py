"""tests/slides/test_slides_templates.py — uploaded-template (A7) pins.

Executable specification for lib/slides/templates.py and the design
stage's template branch: an uploaded brand deck's reverse-engineered
tokens replace the curated theme (geometry included), registry-external
fonts are substituted and REPORTED (never silent — the export gate embeds
registry faces only), and template reads stay owner-fenced.
"""

from __future__ import annotations

import os

import pytest

pytestmark = pytest.mark.unit

from lib.slides.templates import (_normalise_fonts, import_template,
                                  load_template_theme)

TEST_OWNER_USER_ID = 1


def _tokens(**families):
    """e.g. _tokens(title='MiSans', body='Noto Serif SC')."""
    styles = {}
    for role, family in families.items():
        styles[role] = {'fontFamily': family}
    return {'colors': {'bg': '#FFFFFF', 'ink': '#111111',
                       'primary': '#16283C'},
            'textStyles': styles}


# ── font substitution (registry-external faces) ───────────

def test_normalise_fonts_keeps_registry_faces():
    tokens = _tokens(title='MiSans', body='Noto Serif SC')
    assert _normalise_fonts(tokens) == {}
    assert tokens['textStyles']['title']['fontFamily'] == 'MiSans'


def test_normalise_fonts_substitutes_and_reports_unknown():
    tokens = _tokens(title='Microsoft YaHei', body='Songti SC')
    subs = _normalise_fonts(tokens)
    assert subs == {'Microsoft YaHei': 'MiSans',
                    'Songti SC': 'Noto Serif SC'}
    assert tokens['textStyles']['title']['fontFamily'] == 'MiSans'
    assert tokens['textStyles']['body']['fontFamily'] == 'Noto Serif SC'


# ── template.json read fencing ────────────────────────────

def _write_template_record(root, template_id='tpl_abc123', *,
                           owner=TEST_OWNER_USER_ID):
    import json
    dest = os.path.join(str(root), template_id)
    os.makedirs(dest, exist_ok=True)
    record = {
        'template_id': template_id, 'owner_user_id': owner,
        'title': '品牌模板',
        'theme': _tokens(title='MiSans', body='MiSans'),
        'size': [960, 540], 'source_pages': 7,
        'font_substitutions': {'Microsoft YaHei': 'MiSans'},
        'created_at': 1.0,
    }
    with open(os.path.join(dest, 'template.json'), 'w',
              encoding='utf-8') as f:
        json.dump(record, f, ensure_ascii=False)
    return dest


def test_load_template_theme_owner_fenced(tmp_path, monkeypatch):
    monkeypatch.setattr('lib.slides.templates.templates_root',
                        lambda: str(tmp_path))
    _write_template_record(tmp_path)
    hit = load_template_theme('tpl_abc123', owner_user_id=TEST_OWNER_USER_ID)
    assert hit is not None
    assert hit['size'] == [960, 540]
    assert hit['font_substitutions'] == {'Microsoft YaHei': 'MiSans'}
    foreign = load_template_theme('tpl_abc123', owner_user_id=999)
    assert foreign is None
    assert load_template_theme('../etc', owner_user_id=1) is None
    assert load_template_theme('tpl_ghost', owner_user_id=1) is None


# ── import end-to-end ─────────────────────────────────────

def _minimal_pptx(tmp_path) -> str:
    pytest.importorskip('pptx')
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4),
                                   Inches(1))
    box.text_frame.text = '品牌标题'
    path = str(tmp_path / 'brand.pptx')
    prs.save(path)
    return path


def test_import_template_end_to_end(tmp_path, monkeypatch):
    monkeypatch.setattr('lib.slides.templates.templates_root',
                        lambda: str(tmp_path / 'templates'))
    pptx_path = _minimal_pptx(tmp_path)
    result = import_template(pptx_path, title='品牌模板',
                             owner_user_id=TEST_OWNER_USER_ID)
    assert result['template_id'].startswith('tpl_')
    assert result['source_pages'] == 1
    assert set(result['theme']['colors']) == {
        'bg', 'ink', 'primary', 'accent', 'muted', 'hairline'}
    # Re-readable through the fenced path.
    hit = load_template_theme(result['template_id'],
                              owner_user_id=TEST_OWNER_USER_ID)
    assert hit is not None and hit['title'] == '品牌模板'


def test_import_template_requires_owner(tmp_path, monkeypatch):
    monkeypatch.setattr('lib.slides.templates.templates_root',
                        lambda: str(tmp_path / 'templates'))
    pptx_path = _minimal_pptx(tmp_path)
    with pytest.raises(ValueError, match='owner'):
        import_template(pptx_path, owner_user_id=0)


# ── design stage template branch ──────────────────────────

def test_design_stage_binds_template_tokens(tmp_path, monkeypatch):
    monkeypatch.setattr('lib.slides.templates.templates_root',
                        lambda: str(tmp_path))
    _write_template_record(tmp_path)
    monkeypatch.setattr('lib.design_sys.fonts.ensure_font',
                        lambda *a, **k: True)
    import lib.slides.recipe as recipe
    ctx = {'artifacts': {'outline': {'scenario': 'business-plan',
                                     'theme_id': ''}},
           'template_id': 'tpl_abc123', 'owner_user_id': 1,
           'size': (1280, 720)}
    design = recipe._run_design(ctx)
    assert design['theme_id'] == 'template:tpl_abc123'
    assert design['scenario'] == 'business-plan'      # content scenario kept
    assert design['theme_tokens']['colors']['primary'] == '#16283C'
    assert ctx['size'] == (960, 540)                  # template geometry wins
    assert design['font_substitutions'] == {'Microsoft YaHei': 'MiSans'}


def test_design_stage_unknown_template_falls_back(tmp_path, monkeypatch):
    monkeypatch.setattr('lib.slides.templates.templates_root',
                        lambda: str(tmp_path))
    monkeypatch.setattr('lib.design_sys.fonts.ensure_font',
                        lambda *a, **k: True)
    import lib.slides.recipe as recipe
    ctx = {'artifacts': {'outline': {'scenario': 'tech-engineering',
                                     'theme_id': 'paper-engineer'}},
           'template_id': 'tpl_ghost', 'owner_user_id': 1,
           'size': (1280, 720)}
    design = recipe._run_design(ctx)
    assert design['theme_id'] == 'paper-engineer'
    assert ctx['size'] == (1280, 720)


# ── HTTP upload ───────────────────────────────────────────

def test_import_template_endpoint(flask_client, monkeypatch, tmp_path):
    monkeypatch.setattr('lib.slides.templates.templates_root',
                        lambda: str(tmp_path / 'templates'))
    pptx_path = _minimal_pptx(tmp_path)
    with open(pptx_path, 'rb') as f:
        payload = f.read()
    import io
    from werkzeug.datastructures import FileStorage
    upload = FileStorage(stream=io.BytesIO(payload), filename='brand.pptx',
                         content_type='application/vnd.openxmlformats-'
                                      'officedocument.presentationml.'
                                      'presentation')
    r = flask_client.post('/api/v4/slides/import-template',
                          form={'title': '品牌模板'},
                          files={'file': upload})
    assert r.status_code == 200
    body = r.get_json()
    assert body['ok'] and body['template_id'].startswith('tpl_')

    bad_upload = FileStorage(stream=io.BytesIO(b'not a pptx'),
                             filename='notes.txt',
                             content_type='text/plain')
    bad = flask_client.post('/api/v4/slides/import-template',
                            form={}, files={'file': bad_upload})
    assert bad.status_code == 400
